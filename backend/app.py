from flask import Flask, request, jsonify, send_file, send_from_directory
import os
import json
import glob
import base64
import pickle
import smtplib
import urllib.request
from datetime import datetime
from collections import Counter
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders

import cv2
import numpy as np
import mediapipe as mp
from mediapipe.tasks.python import BaseOptions
from mediapipe.tasks.python.vision import FaceLandmarker, FaceLandmarkerOptions, RunningMode
from flask_cors import CORS
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── Configuration ──────────────────────────────────────────────────────
BACKEND_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BACKEND_DIR, 'Data')
ATTENDANCE_DIR = os.path.join(BACKEND_DIR, 'Attendance')
CASCADE_PATH = os.path.join(DATA_DIR, 'haarcascade_frontalface_default.xml')
GROUPS_PATH = os.path.join(DATA_DIR, 'groups.json')
ENCODINGS_PATH = os.path.join(DATA_DIR, 'face_encodings.pkl')
LABELS_PATH = os.path.join(DATA_DIR, 'names.pkl')

os.makedirs(ATTENDANCE_DIR, exist_ok=True)
os.makedirs(DATA_DIR, exist_ok=True)

DEBUG = os.environ.get('FLASK_DEBUG', 'false').lower() in ('1', 'true', 'yes')
PORT = int(os.environ.get('PORT', 5000))
CORS_ORIGINS = os.environ.get('CORS_ORIGINS', '*')

FRONTEND_DIST = os.path.join(BACKEND_DIR, '..', 'Fronend', 'dist')

app = Flask(__name__, static_folder=None)
CORS(app, origins=CORS_ORIGINS.split(','))
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50 MB

DEEPFACE_MODEL = 'SFace'
MATCH_THRESHOLD = 0.40
FACE_LANDMARKER_MODEL = os.path.join(DATA_DIR, 'face_landmarker.task')

LANDMARKER_MODEL_URL = (
    'https://storage.googleapis.com/mediapipe-models/face_landmarker/face_landmarker/float16/1/face_landmarker.task'
)


def _ensure_models():
    """Download required model files if they are missing."""
    if not os.path.isfile(FACE_LANDMARKER_MODEL):
        print("Downloading face_landmarker.task ...")
        urllib.request.urlretrieve(LANDMARKER_MODEL_URL, FACE_LANDMARKER_MODEL)
        print(f"Downloaded -> {FACE_LANDMARKER_MODEL}")


_ensure_models()

LEFT_EYE_IDX = [33, 160, 158, 133, 153, 144]
RIGHT_EYE_IDX = [362, 385, 387, 263, 373, 380]
EAR_THRESHOLD = 0.21
MIN_BLINKS = 2

_deepface_module = None


def _deepface():
    global _deepface_module
    if _deepface_module is None:
        os.environ['TF_CPP_MIN_LOG_LEVEL'] = '3'
        from deepface import DeepFace
        _deepface_module = DeepFace
    return _deepface_module


def _prewarm_deepface():
    """Load the DeepFace model at startup so the first request isn't slow."""
    import sys, io
    if sys.platform == 'win32':
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
        sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')
    try:
        print("Pre-warming DeepFace model...")
        dummy = np.zeros((100, 100, 3), dtype=np.uint8)
        _deepface().represent(
            img_path=dummy,
            model_name=DEEPFACE_MODEL,
            detector_backend='opencv',
            enforce_detection=False,
        )
        print("DeepFace model ready.")
    except Exception:
        print("DeepFace prewarm: model will load on first request.")


_prewarm_deepface()


# ── Data helpers ───────────────────────────────────────────────────────

def load_groups():
    if os.path.isfile(GROUPS_PATH):
        with open(GROUPS_PATH, 'r') as f:
            return json.load(f)
    return {}


def save_groups(groups):
    with open(GROUPS_PATH, 'w') as f:
        json.dump(groups, f, indent=2)


def decode_frame(frame_b64):
    if ',' in frame_b64:
        frame_b64 = frame_b64.split(',')[1]
    img_bytes = base64.b64decode(frame_b64)
    nparr = np.frombuffer(img_bytes, np.uint8)
    return cv2.imdecode(nparr, cv2.IMREAD_COLOR)


def load_face_db():
    if not os.path.isfile(ENCODINGS_PATH) or not os.path.isfile(LABELS_PATH):
        return None, None, "No registered faces found. Please register first."
    with open(ENCODINGS_PATH, 'rb') as f:
        encodings = pickle.load(f)
    with open(LABELS_PATH, 'rb') as f:
        labels = pickle.load(f)
    if len(encodings) != len(labels):
        return None, None, f"Data mismatch: {len(encodings)} encodings vs {len(labels)} labels. Re-register."
    return np.array(encodings), labels, None


# ── Face recognition helpers ──────────────────────────────────────────

def cosine_sim(a, b):
    a = np.asarray(a, dtype=np.float64)
    b = np.asarray(b, dtype=np.float64)
    return float(np.dot(a, b) / (np.linalg.norm(a) * np.linalg.norm(b) + 1e-10))


def find_best_match(embedding, stored_encodings, stored_labels):
    if stored_encodings is None or len(stored_encodings) == 0:
        return "Unknown", "", 0.0
    sims = np.array([cosine_sim(embedding, s) for s in stored_encodings])
    best_idx = int(np.argmax(sims))
    best_sim = float(sims[best_idx])
    if best_sim >= MATCH_THRESHOLD:
        label = stored_labels[best_idx]
        parts = label.split(',', 1)
        name = parts[0]
        uid = parts[1] if len(parts) > 1 else ""
        return name, uid, best_sim
    return "Unknown", "", best_sim


def extract_faces_from_image(img):
    """Return list of {embedding, facial_area, face_confidence} using DeepFace."""
    DeepFace = _deepface()
    try:
        results = DeepFace.represent(
            img_path=img,
            model_name=DEEPFACE_MODEL,
            detector_backend='opencv',
            enforce_detection=False,
        )
        return [r for r in results if r.get('face_confidence', 0) > 0.5]
    except Exception as e:
        print(f"DeepFace extraction error: {e}")
        return []


# ── Liveness detection helpers ────────────────────────────────────────

def compute_ear(landmarks, h, w):
    """Compute average Eye Aspect Ratio from a list of NormalizedLandmark."""
    def _eye_ear(indices):
        pts = [(landmarks[i].x * w, landmarks[i].y * h) for i in indices]
        v1 = np.linalg.norm(np.array(pts[1]) - np.array(pts[5]))
        v2 = np.linalg.norm(np.array(pts[2]) - np.array(pts[4]))
        horiz = np.linalg.norm(np.array(pts[0]) - np.array(pts[3]))
        return (v1 + v2) / (2.0 * horiz + 1e-6)
    left = _eye_ear(LEFT_EYE_IDX)
    right = _eye_ear(RIGHT_EYE_IDX)
    return (left + right) / 2.0


def _get_face_landmarker():
    """Create a FaceLandmarker using the Tasks API."""
    options = FaceLandmarkerOptions(
        base_options=BaseOptions(model_asset_path=FACE_LANDMARKER_MODEL),
        running_mode=RunningMode.IMAGE,
        num_faces=1,
        min_face_detection_confidence=0.5,
        min_face_presence_confidence=0.5,
    )
    return FaceLandmarker.create_from_options(options)


def count_blinks(ear_values):
    blinks = 0
    below = 0
    for ear in ear_values:
        if ear < EAR_THRESHOLD:
            below += 1
        else:
            if below >= 2:
                blinks += 1
            below = 0
    return blinks


# ── Attendance Excel helpers ──────────────────────────────────────────

def get_group_dir(group):
    group_dir = os.path.join(ATTENDANCE_DIR, group)
    os.makedirs(group_dir, exist_ok=True)
    return group_dir


def build_attendance_excel(group, date_str, present_labels, time_str):
    groups = load_groups()
    students = groups.get(group, [])

    group_dir = get_group_dir(group)
    filepath = os.path.join(group_dir, f"{group}_{date_str}.xlsx")

    already_present = {}
    if os.path.isfile(filepath):
        wb = load_workbook(filepath)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=False):
            label = row[2].value
            if row[3].value == "Present":
                already_present[label] = row[4].value

    present_set = set(already_present.keys())
    new_count = 0
    for label in present_labels:
        if label not in present_set:
            present_set.add(label)
            already_present[label] = time_str
            new_count += 1

    wb = Workbook()
    ws = wb.active
    ws.title = f"{group} - {date_str}"

    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    present_fill = PatternFill(start_color="d4edda", end_color="d4edda", fill_type="solid")
    absent_fill = PatternFill(start_color="f8d7da", end_color="f8d7da", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    headers = ["S.No", "Name", "UID", "Status", "Time"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    row_num = 2
    for i, student in enumerate(students, 1):
        label = f"{student['name']},{student['uid']}"
        is_present = label in present_set
        status = "Present" if is_present else "Absent"
        marked_time = already_present.get(label, "") if is_present else ""
        fill = present_fill if is_present else absent_fill

        values = [i, student['name'], student['uid'], status, marked_time]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center')
        row_num += 1

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12

    present_count = len(present_set)
    total = len(students)
    ws.cell(row=row_num + 1, column=1, value="Summary:").font = Font(bold=True)
    ws.cell(row=row_num + 1, column=2, value=f"Present: {present_count}/{total}")
    ws.cell(row=row_num + 1, column=3, value=f"Absent: {total - present_count}/{total}")

    wb.save(filepath)
    return filepath, new_count, present_count, total


def read_excel(filepath):
    wb = load_workbook(filepath, read_only=True)
    ws = wb.active
    records = []
    headers = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(h) for h in row]
            continue
        if row[0] is None or not str(row[0]).isdigit():
            break
        record = {}
        for j, header in enumerate(headers):
            record[header] = str(row[j]) if row[j] is not None else ""
        records.append(record)
    wb.close()
    return records


# ══════════════════════════════════════════════════════════════════════
#  ENDPOINTS
# ══════════════════════════════════════════════════════════════════════

@app.route('/register-face', methods=['POST'])
def register_face():
    data = request.get_json()
    name = data.get('name')
    uid = data.get('uid')
    group = data.get('group')
    frames = data.get('frames', [])

    if not name or not uid:
        return jsonify({"error": "Name and UID are required"}), 400
    if not group:
        return jsonify({"error": "Group is required"}), 400
    if len(frames) < 5:
        return jsonify({"error": "At least 5 face frames are required"}), 400

    embeddings = []
    for frame_b64 in frames:
        try:
            img = decode_frame(frame_b64)
            if img is None:
                continue
            faces = extract_faces_from_image(img)
            if faces:
                embeddings.append(faces[0]['embedding'])
        except Exception as e:
            print(f"Frame processing error: {e}")
            continue

    if len(embeddings) < 5:
        return jsonify({
            "error": f"Only {len(embeddings)} face embeddings extracted. Need at least 5. Try better lighting."
        }), 400

    label = f"{name},{uid}"
    count = len(embeddings)

    if os.path.isfile(LABELS_PATH):
        with open(LABELS_PATH, 'rb') as f:
            labels = pickle.load(f)
        labels = labels + [label] * count
    else:
        labels = [label] * count
    with open(LABELS_PATH, 'wb') as f:
        pickle.dump(labels, f)

    if os.path.isfile(ENCODINGS_PATH):
        with open(ENCODINGS_PATH, 'rb') as f:
            existing = pickle.load(f)
        all_enc = existing + embeddings
    else:
        all_enc = embeddings
    with open(ENCODINGS_PATH, 'wb') as f:
        pickle.dump(all_enc, f)

    groups = load_groups()
    if group not in groups:
        groups[group] = []
    if not any(s['uid'] == uid for s in groups[group]):
        groups[group].append({"name": name, "uid": uid})
    save_groups(groups)

    return jsonify({
        "message": "Registration successful",
        "faces_captured": count,
        "name": name,
        "uid": uid,
        "group": group,
    }), 200


@app.route('/take-attendance', methods=['POST'])
def take_attendance():
    data = request.get_json()
    group = data.get('group')
    frames = data.get('frames', [])

    if not group:
        return jsonify({"error": "Group is required"}), 400
    if len(frames) < 3:
        return jsonify({"error": "At least 3 frames are required"}), 400

    groups = load_groups()
    if group not in groups or len(groups[group]) == 0:
        return jsonify({"error": f"No students registered in {group}. Register students first."}), 400

    stored_enc, stored_labels, err = load_face_db()
    if stored_enc is None:
        return jsonify({"error": err}), 400

    all_predictions = []
    for frame_b64 in frames:
        try:
            img = decode_frame(frame_b64)
            if img is None:
                continue
            faces = extract_faces_from_image(img)
            for face in faces:
                name, uid, conf = find_best_match(face['embedding'], stored_enc, stored_labels)
                if name != "Unknown":
                    all_predictions.append(f"{name},{uid}")
        except Exception as e:
            print(f"Recognition error: {e}")
            continue

    if not all_predictions:
        return jsonify({"error": "No faces recognized. Ensure someone registered is in front of the camera."}), 400

    counts = Counter(all_predictions)
    now = datetime.now()
    date_str = now.strftime("%d-%m-%Y")
    time_str = now.strftime("%H:%M:%S")

    present_labels = [label for label, _ in counts.most_common()]
    filepath, new_count, present_count, total = build_attendance_excel(
        group, date_str, present_labels, time_str,
    )

    recognized = []
    for label, cnt in counts.most_common():
        name_part = label.split(',')[0] if ',' in label else label
        recognized.append({"name": name_part, "label": label, "count": cnt})

    return jsonify({
        "recognized": recognized,
        "new_entries": new_count,
        "already_marked": len(recognized) - new_count,
        "present": present_count,
        "total_students": total,
        "group": group,
        "date": date_str,
        "time": time_str,
        "file": os.path.basename(filepath),
    }), 200


@app.route('/recognize-frame', methods=['POST'])
def recognize_frame():
    """Single-frame recognition returning face bounding boxes + names."""
    data = request.get_json()
    frame_b64 = data.get('frame', '')
    group = data.get('group', '')

    if not frame_b64:
        return jsonify({"faces": []}), 200

    stored_enc, stored_labels, _ = load_face_db()

    img = decode_frame(frame_b64)
    if img is None:
        return jsonify({"faces": []}), 200

    detected = extract_faces_from_image(img)
    faces_out = []
    for face in detected:
        embedding = face['embedding']
        area = face.get('facial_area', {})
        name, uid, conf = find_best_match(embedding, stored_enc, stored_labels)
        faces_out.append({
            "name": name,
            "uid": uid,
            "x": area.get('x', 0),
            "y": area.get('y', 0),
            "w": area.get('w', 0),
            "h": area.get('h', 0),
            "confidence": round(conf, 2),
        })

    return jsonify({"faces": faces_out}), 200


@app.route('/liveness-check', methods=['POST'])
def liveness_check():
    """Accept a sequence of frames and check for eye blinks via mediapipe Tasks API."""
    data = request.get_json()
    frames = data.get('frames', [])

    if len(frames) < 10:
        return jsonify({"alive": False, "blinks_detected": 0, "error": "Need at least 10 frames"}), 400

    landmarker = _get_face_landmarker()

    ear_values = []
    for frame_b64 in frames:
        try:
            img = decode_frame(frame_b64)
            if img is None:
                continue
            h, w = img.shape[:2]
            rgb = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
            mp_image = mp.Image(image_format=mp.ImageFormat.SRGB, data=rgb)
            result = landmarker.detect(mp_image)
            if result.face_landmarks and len(result.face_landmarks) > 0:
                landmarks = result.face_landmarks[0]
                ear = compute_ear(landmarks, h, w)
                ear_values.append(ear)
        except Exception as e:
            print(f"Liveness frame error: {e}")
            continue

    landmarker.close()

    if len(ear_values) < 5:
        return jsonify({"alive": False, "blinks_detected": 0, "error": "Could not detect face landmarks"}), 400

    blinks = count_blinks(ear_values)
    alive = blinks >= MIN_BLINKS

    return jsonify({
        "alive": alive,
        "blinks_detected": blinks,
        "frames_processed": len(ear_values),
    }), 200


@app.route('/view-attendance', methods=['GET'])
def view_attendance():
    group = request.args.get('group')
    date_param = request.args.get('date')

    if not group:
        return jsonify({"records": [], "dates": [], "message": "Select a group to view attendance"}), 200

    group_dir = os.path.join(ATTENDANCE_DIR, group)
    if not os.path.isdir(group_dir):
        return jsonify({"records": [], "dates": [], "group": group, "message": f"No attendance records for {group}"}), 200

    pattern = os.path.join(group_dir, f"{group}_*.xlsx")
    files = glob.glob(pattern)
    if not files:
        return jsonify({"records": [], "dates": [], "group": group, "message": f"No attendance records for {group}"}), 200

    files.sort(key=os.path.getmtime, reverse=True)

    available_dates = []
    for f in files:
        fn = os.path.basename(f)
        d = fn.replace(f"{group}_", "").replace(".xlsx", "")
        available_dates.append(d)

    target_date = date_param if date_param else available_dates[0]
    target_file = os.path.join(group_dir, f"{group}_{target_date}.xlsx")

    if not os.path.isfile(target_file):
        return jsonify({
            "records": [], "date": target_date, "dates": available_dates,
            "group": group, "message": "No attendance for this date",
        }), 200

    records = read_excel(target_file)
    return jsonify({"records": records, "date": target_date, "dates": available_dates, "group": group}), 200


@app.route('/download-attendance', methods=['GET'])
def download_attendance():
    group = request.args.get('group')
    date = request.args.get('date')
    if not group or not date:
        return jsonify({"error": "Group and date are required"}), 400
    filepath = os.path.join(ATTENDANCE_DIR, group, f"{group}_{date}.xlsx")
    if not os.path.isfile(filepath):
        return jsonify({"error": "File not found"}), 404
    return send_file(filepath, as_attachment=True, download_name=f"{group}_{date}.xlsx")


@app.route('/groups', methods=['GET'])
def list_groups():
    """Return all available group names from groups.json."""
    groups = load_groups()
    return jsonify({"groups": sorted(groups.keys())}), 200


@app.route('/group-students', methods=['GET'])
def group_students():
    group = request.args.get('group')
    if not group:
        return jsonify({"students": []}), 200
    groups = load_groups()
    students = groups.get(group, [])

    registered_uids = set()
    if os.path.isfile(LABELS_PATH):
        with open(LABELS_PATH, 'rb') as f:
            labels = pickle.load(f)
        for label in labels:
            parts = label.split(',', 1)
            if len(parts) > 1:
                registered_uids.add(parts[1])

    enriched = []
    for s in students:
        enriched.append({
            "name": s["name"],
            "uid": s["uid"],
            "face_registered": s["uid"] in registered_uids,
        })

    registered_count = sum(1 for s in enriched if s["face_registered"])
    return jsonify({
        "students": enriched,
        "group": group,
        "count": len(enriched),
        "registered_count": registered_count,
        "unregistered_count": len(enriched) - registered_count,
    }), 200


@app.route('/analytics', methods=['GET'])
def analytics():
    group = request.args.get('group')
    if not group:
        return jsonify({"error": "Group is required"}), 400

    group_dir = os.path.join(ATTENDANCE_DIR, group)
    groups_data = load_groups()
    students = groups_data.get(group, [])
    total_students = len(students)

    empty = {
        "group": group, "total_students": total_students, "total_classes": 0,
        "dates": [], "daily_stats": [], "student_stats": [], "at_risk_count": 0,
    }

    if not os.path.isdir(group_dir):
        return jsonify(empty), 200

    pattern = os.path.join(group_dir, f"{group}_*.xlsx")
    files = sorted(glob.glob(pattern), key=os.path.getmtime)
    if not files:
        return jsonify(empty), 200

    student_present_count = {s['uid']: 0 for s in students}
    daily_stats = []
    dates = []

    for filepath in files:
        fn = os.path.basename(filepath)
        date_str = fn.replace(f"{group}_", "").replace(".xlsx", "")
        dates.append(date_str)

        records = read_excel(filepath)
        present = 0
        for r in records:
            uid = r.get('UID', '')
            if r.get('Status') == 'Present':
                present += 1
                if uid in student_present_count:
                    student_present_count[uid] += 1

        pct = (present / total_students * 100) if total_students > 0 else 0
        daily_stats.append({
            "date": date_str,
            "present": present,
            "absent": total_students - present,
            "percentage": round(pct, 1),
        })

    total_classes = len(files)
    student_stats = []
    at_risk_count = 0
    for s in students:
        p = student_present_count.get(s['uid'], 0)
        pct = (p / total_classes * 100) if total_classes > 0 else 0
        at_risk = pct < 75.0 and total_classes > 0
        if at_risk:
            at_risk_count += 1
        student_stats.append({
            "name": s['name'], "uid": s['uid'],
            "total_present": p, "total_classes": total_classes,
            "percentage": round(pct, 1), "at_risk": at_risk,
        })

    student_stats.sort(key=lambda x: x['percentage'])

    return jsonify({
        "group": group, "total_students": total_students,
        "total_classes": total_classes, "dates": dates,
        "daily_stats": daily_stats, "student_stats": student_stats,
        "at_risk_count": at_risk_count,
    }), 200


@app.route('/send-report', methods=['POST'])
def send_report():
    data = request.get_json()
    group = data.get('group')
    date = data.get('date')
    email_to = data.get('email')

    if not group or not date or not email_to:
        return jsonify({"error": "group, date, and email are required"}), 400

    filepath = os.path.join(ATTENDANCE_DIR, group, f"{group}_{date}.xlsx")
    if not os.path.isfile(filepath):
        return jsonify({"error": "Attendance file not found"}), 404

    smtp_host = os.environ.get('SMTP_HOST', 'smtp.gmail.com')
    smtp_port = int(os.environ.get('SMTP_PORT', '587'))
    smtp_user = os.environ.get('SMTP_USER', '')
    smtp_pass = os.environ.get('SMTP_PASS', '')

    if not smtp_user or not smtp_pass:
        return jsonify({"error": "Email not configured. Set SMTP_USER and SMTP_PASS environment variables."}), 500

    try:
        msg = MIMEMultipart()
        msg['From'] = smtp_user
        msg['To'] = email_to
        msg['Subject'] = f"Attendance Report - {group} - {date}"

        body = (
            f"Hi,\n\n"
            f"Please find attached the attendance report for {group} on {date}.\n\n"
            f"Generated by Drishtikon Attendance System.\n"
        )
        msg.attach(MIMEText(body, 'plain'))

        with open(filepath, 'rb') as f:
            part = MIMEBase('application', 'octet-stream')
            part.set_payload(f.read())
        encoders.encode_base64(part)
        part.add_header('Content-Disposition', f'attachment; filename="{group}_{date}.xlsx"')
        msg.attach(part)

        with smtplib.SMTP(smtp_host, smtp_port) as server:
            server.starttls()
            server.login(smtp_user, smtp_pass)
            server.send_message(msg)

        return jsonify({"message": f"Report sent to {email_to}"}), 200
    except Exception as e:
        return jsonify({"error": f"Failed to send email: {str(e)}"}), 500


# ── Serve frontend in production ───────────────────────────────────────

@app.route('/', defaults={'path': ''})
@app.route('/<path:path>')
def serve_frontend(path):
    """Serve the React build. API routes are registered first so they take priority."""
    dist = os.path.abspath(FRONTEND_DIST)
    if path and os.path.isfile(os.path.join(dist, path)):
        return send_from_directory(dist, path)
    index = os.path.join(dist, 'index.html')
    if os.path.isfile(index):
        return send_from_directory(dist, 'index.html')
    return jsonify({"error": "Frontend not built. Run: cd Fronend && npm run build"}), 404


if __name__ == '__main__':
    if DEBUG:
        app.run(host='0.0.0.0', port=PORT, debug=True)
    else:
        from waitress import serve
        print(f"Starting production server on port {PORT}")
        serve(app, host='0.0.0.0', port=PORT, threads=4)
